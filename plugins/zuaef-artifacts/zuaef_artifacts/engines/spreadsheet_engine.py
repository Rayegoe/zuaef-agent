"""Spreadsheet engine: declarative workbook spec → .xlsx via openpyxl,
bounded revisions, formula preservation, and optional LibreOffice
recalculation.

Formula policy (spec pack 07): derived values that express business logic
stay formulas so the user can keep editing inputs. openpyxl does not evaluate
formulas; when LibreOffice is available a bounded recalculation pass embeds
cached results, otherwise the limitation is reported — never fabricated
values.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Annotated, Any, Literal

from pydantic import BaseModel, Field, TypeAdapter, ValidationError

from .errors import UnsupportedOperation

MAX_SHEETS = 20
MAX_BLOCKS_PER_SHEET = 100
MAX_MATRIX_ROWS = 500
MAX_MATRIX_COLS = 60
MAX_CELL_CHARS = 5_000
MAX_TABLE_ROWS = 500

_CELL_REF = re.compile(r"^[A-Za-z]{1,3}[1-9][0-9]*$")
_RANGE_REF = re.compile(r"^([A-Za-z]{1,3}[1-9][0-9]*):([A-Za-z]{1,3}[1-9][0-9]*)$")


def _validate_range(ref: str) -> None:
    """A block range is a single cell ('B5') or a rectangle ('A1:C5')."""
    if not (_CELL_REF.fullmatch(ref) or _RANGE_REF.fullmatch(ref)):
        raise ValueError(f"invalid range (use 'B5' or 'A1:C5'): {ref!r}")


def _range_bounds(ref: str) -> tuple[int, int, int, int]:
    if ":" in ref:
        start, end = ref.split(":")
        (row1, col1), (row2, col2) = _cell_to_row_col(start), _cell_to_row_col(end)
        return min(row1, row2), min(col1, col2), max(row1, row2), max(col1, col2)
    row, col = _cell_to_row_col(ref)
    return row, col, row, col


class MatrixBlock(BaseModel):
    """Write a 2-D block of values/formulas at start_cell (e.g. 'A1').
    Strings starting with '=' are written as formulas."""

    kind: Literal["matrix"] = "matrix"
    start_cell: str
    values: list[list[Any]]

    def model_post_init(self, __context: Any, /) -> None:
        if not _CELL_REF.fullmatch(self.start_cell):
            raise ValueError(f"invalid start_cell: {self.start_cell!r}")
        if not self.values or len(self.values) > MAX_MATRIX_ROWS:
            raise ValueError(f"matrix needs 1..{MAX_MATRIX_ROWS} rows")
        width = len(self.values[0])
        if width < 1 or width > MAX_MATRIX_COLS:
            raise ValueError(f"matrix needs 1..{MAX_MATRIX_COLS} columns")
        for row in self.values:
            if len(row) != width:
                raise ValueError("every matrix row must have the same width")
            for value in row:
                if isinstance(value, str) and len(value) > MAX_CELL_CHARS:
                    raise ValueError("matrix cell text too long")


class FormatBlock(BaseModel):
    """Apply number/font/fill/border/alignment formatting to a range."""

    kind: Literal["format"] = "format"
    range: str
    number_format: str | None = Field(default=None, max_length=60)
    bold: bool | None = None
    font_size: float | None = Field(default=None, gt=0, le=72)
    font_color: str | None = Field(default=None, description="RRGGBB")
    fill_color: str | None = Field(default=None, description="RRGGBB")
    border: bool | None = None
    halign: Literal["left", "center", "right"] | None = None
    wrap: bool | None = None

    def model_post_init(self, __context: Any, /) -> None:
        _validate_range(self.range)
        for color in (self.font_color, self.fill_color):
            if color is not None and not re.fullmatch(r"[0-9A-Fa-f]{6}", color):
                raise ValueError(f"colors must be RRGGBB hex: {color!r}")


class WidthsBlock(BaseModel):
    """Set column widths (letters → width in Excel units)."""

    kind: Literal["widths"] = "widths"
    widths: dict[str, float]

    def model_post_init(self, __context: Any, /) -> None:
        if not self.widths or len(self.widths) > 60:
            raise ValueError("widths need 1..60 columns")
        for column in self.widths:
            if not re.fullmatch(r"[A-Za-z]{1,3}", column):
                raise ValueError(f"invalid column letter: {column!r}")
        for width in self.widths.values():
            if not 1 <= width <= 200:
                raise ValueError("column widths must be in [1, 200]")


class ValidationBlock(BaseModel):
    """Restrict a range to a list of allowed values."""

    kind: Literal["validation"] = "validation"
    range: str
    allowed: list[str]

    def model_post_init(self, __context: Any, /) -> None:
        _validate_range(self.range)
        if not self.allowed or len(self.allowed) > 50:
            raise ValueError("validation needs 1..50 allowed values")
        total = sum(len(value) for value in self.allowed)
        if total > 800:
            raise ValueError("validation allowed list too long")


class ConditionalFormatBlock(BaseModel):
    """Cell-is conditional formatting rule over a range."""

    kind: Literal["conditional_format"] = "conditional_format"
    range: str
    operator: Literal["greater_than", "less_than", "equal", "between"]
    formula: list[str] = Field(min_length=1, max_length=2)
    fill_color: str | None = Field(default=None, description="RRGGBB")

    def model_post_init(self, __context: Any, /) -> None:
        _validate_range(self.range)
        if self.operator == "between" and len(self.formula) != 2:
            raise ValueError("between needs exactly two formulas")
        if self.fill_color is not None and not re.fullmatch(
            r"[0-9A-Fa-f]{6}", self.fill_color
        ):
            raise ValueError(f"fill_color must be RRGGBB hex: {self.fill_color!r}")


class TableBlock(BaseModel):
    """Register a structured Excel table over a range (header row included)."""

    kind: Literal["table"] = "table"
    range: str
    name: str = Field(default="Table1", max_length=40)
    style: str = Field(default="TableStyleMedium9", max_length=40)

    def model_post_init(self, __context: Any, /) -> None:
        _validate_range(self.range)
        if not re.fullmatch(r"[A-Za-z0-9_]+", self.name):
            raise ValueError("table name must be alphanumeric/underscore")


class ChartBlock(BaseModel):
    """Add a basic chart anchored at a cell. data_range includes the header
    row with series names; categories_range points at the label column."""

    kind: Literal["chart"] = "chart"
    chart_type: Literal["bar", "line", "pie"] = "bar"
    anchor: str
    data_range: str
    categories_range: str | None = None
    title: str | None = Field(default=None, max_length=120)

    def model_post_init(self, __context: Any, /) -> None:
        for ref in (self.anchor,):
            if not _CELL_REF.fullmatch(ref):
                raise ValueError(f"invalid anchor cell: {ref!r}")
        for ref in (self.data_range, self.categories_range):
            if ref is not None and not _RANGE_REF.fullmatch(ref):
                raise ValueError(f"invalid range (use 'B2:D5'): {ref!r}")


SheetBlock = Annotated[
    MatrixBlock | FormatBlock | WidthsBlock | ValidationBlock | ConditionalFormatBlock | TableBlock | ChartBlock,
    Field(discriminator="kind"),
]


class SheetSpec(BaseModel):
    name: str = Field(max_length=31)
    blocks: list[SheetBlock]

    def model_post_init(self, __context: Any, /) -> None:
        if not re.fullmatch(r"[^[\]*?/\\:]+", self.name):
            raise ValueError(f"invalid sheet name: {self.name!r}")
        if not self.blocks or len(self.blocks) > MAX_BLOCKS_PER_SHEET:
            raise ValueError(f"sheet needs 1..{MAX_BLOCKS_PER_SHEET} blocks")


class WorkbookSpec(BaseModel):
    sheets: list[SheetSpec]

    def model_post_init(self, __context: Any, /) -> None:
        if not self.sheets or len(self.sheets) > MAX_SHEETS:
            raise ValueError(f"workbook needs 1..{MAX_SHEETS} sheets")
        names = [sheet.name for sheet in self.sheets]
        if len(set(names)) != len(names):
            raise ValueError("sheet names must be unique")


class ApplyBlocksOp(BaseModel):
    """Apply sheet blocks to an existing sheet (creation blocks, revisited)."""

    kind: Literal["apply_blocks"] = "apply_blocks"
    sheet: str = Field(max_length=31)
    blocks: list[SheetBlock]

    def model_post_init(self, __context: Any, /) -> None:
        if not self.blocks or len(self.blocks) > 50:
            raise ValueError("apply_blocks needs 1..50 blocks")


class AddSheetOp(BaseModel):
    kind: Literal["add_sheet"] = "add_sheet"
    sheet: SheetSpec


SheetRevisionOp = Annotated[ApplyBlocksOp | AddSheetOp, Field(discriminator="kind")]


def workbook_from_json(data: dict) -> WorkbookSpec:
    try:
        return WorkbookSpec.model_validate(data)
    except ValidationError as exc:
        first = exc.errors()[0]
        location = ".".join(str(part) for part in first.get("loc", ()))
        raise ValueError(
            f"invalid workbook spec at {location or '<root>'}: {first.get('msg')}"
        ) from exc


def sheet_ops_from_json(data: list[dict]) -> list[SheetRevisionOp]:
    try:
        ops = TypeAdapter(list[SheetRevisionOp]).validate_python(data)
    except ValidationError as exc:
        first = exc.errors()[0]
        location = ".".join(str(part) for part in first.get("loc", ()))
        raise ValueError(
            f"invalid sheet operation at {location or '<root>'}: {first.get('msg')}"
        ) from exc
    if not ops:
        raise ValueError("revision has no operations")
    if len(ops) > 30:
        raise ValueError("revision exceeds 30 operations")
    return ops


def _cell_to_row_col(ref: str) -> tuple[int, int]:
    match = re.fullmatch(r"([A-Za-z]{1,3})([1-9][0-9]*)", ref)
    if not match:  # pragma: no cover - models validate first
        raise ValueError(f"invalid cell reference: {ref!r}")
    letters, digits = match.groups()
    column = 0
    for char in letters.upper():
        column = column * 26 + (ord(char) - ord("A") + 1)
    return int(digits), column


def _apply_block(ws, block: Any, workbook) -> None:
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table as XLTable
    from openpyxl.worksheet.table import TableStyleInfo

    kind = block.kind

    if kind == "matrix":
        row, col = _cell_to_row_col(block.start_cell)
        for r_offset, row_values in enumerate(block.values):
            for c_offset, value in enumerate(row_values):
                ws.cell(row=row + r_offset, column=col + c_offset, value=value)
        return

    if kind == "format":
        min_row, min_col, max_row, max_col = _range_bounds(block.range)
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                if block.number_format:
                    cell.number_format = block.number_format
                if block.bold is not None or block.font_size or block.font_color:
                    font_kwargs: dict[str, Any] = {}
                    if block.bold is not None:
                        font_kwargs["bold"] = block.bold
                    if block.font_size:
                        font_kwargs["size"] = block.font_size
                    if block.font_color:
                        font_kwargs["color"] = block.font_color
                    cell.font = Font(**font_kwargs)
                if block.fill_color:
                    cell.fill = PatternFill(
                        start_color=block.fill_color,
                        end_color=block.fill_color,
                        fill_type="solid",
                    )
                if block.border is True:
                    side = Side(style="thin")
                    cell.border = Border(left=side, right=side, top=side, bottom=side)
                if block.halign or block.wrap is not None:
                    alignment_kwargs: dict[str, Any] = {}
                    if block.halign:
                        alignment_kwargs["horizontal"] = block.halign
                    if block.wrap is not None:
                        alignment_kwargs["wrap_text"] = block.wrap
                    cell.alignment = Alignment(**alignment_kwargs)
        return

    if kind == "widths":
        for column, width in block.widths.items():
            ws.column_dimensions[column.upper()].width = width
        return

    if kind == "validation":
        min_row, min_col, max_row, max_col = _range_bounds(block.range)
        joined = ",".join(str(value) for value in block.allowed)
        validation = DataValidation(type="list", formula1=f'"{joined}"', allow_blank=True)
        ws.add_data_validation(validation)
        validation.add(f"{ws.cell(row=min_row, column=min_col).coordinate}:{ws.cell(row=max_row, column=max_col).coordinate}")
        return

    if kind == "conditional_format":
        min_row, min_col, max_row, max_col = _range_bounds(block.range)
        operator_map = {
            "greater_than": "greaterThan",
            "less_than": "lessThan",
            "equal": "equal",
            "between": "between",
        }
        rule = CellIsRule(
            operator=operator_map[block.operator],
            formula=[str(f) for f in block.formula],
            fill=PatternFill(
                start_color=block.fill_color or "FFC7CE",
                end_color=block.fill_color or "FFC7CE",
                fill_type="solid",
            ),
        )
        ws.conditional_formatting.add(
            f"{ws.cell(row=min_row, column=min_col).coordinate}:"
            f"{ws.cell(row=max_row, column=max_col).coordinate}",
            rule,
        )
        return

    if kind == "table":
        table = XLTable(displayName=block.name, ref=block.range)
        table.tableStyleInfo = TableStyleInfo(
            name=block.style, showRowStripes=True
        )
        ws.add_table(table)
        return

    if kind == "chart":
        chart_class = {"bar": BarChart, "line": LineChart, "pie": PieChart}[block.chart_type]
        chart = chart_class()
        chart.title = block.title or None
        min_row, min_col, max_row, max_col = _range_bounds(block.data_range)
        data = Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
        chart.add_data(data, titles_from_data=True)
        if block.categories_range:
            c_min_row, c_min_col, c_max_row, c_max_col = _range_bounds(block.categories_range)
            categories = Reference(
                ws, min_col=c_min_col, min_row=c_min_row, max_col=c_max_col, max_row=c_max_row
            )
            chart.set_categories(categories)
        ws.add_chart(chart, block.anchor)
        return

    raise UnsupportedOperation(f"unknown sheet block: {kind!r}")


def build_workbook(spec: WorkbookSpec) -> Any:
    """Render a WorkbookSpec into an openpyxl Workbook (unsaved)."""
    from openpyxl import Workbook

    workbook = Workbook()
    default_sheet = workbook.active
    for index, sheet_spec in enumerate(spec.sheets):
        if index == 0:
            ws = default_sheet
            ws.title = sheet_spec.name
        else:
            ws = workbook.create_sheet(title=sheet_spec.name)
        for block in sheet_spec.blocks:
            _apply_block(ws, block, workbook)
    return workbook


def resolve_sheet(workbook: Any, name: str) -> Any:
    if name in workbook.sheetnames:
        return workbook[name]
    raise UnsupportedOperation(f"sheet not found: {name!r}")


def apply_sheet_op(workbook: Any, op: Any) -> str:
    if op.kind == "apply_blocks":
        ws = resolve_sheet(workbook, op.sheet)
        for block in op.blocks:
            _apply_block(ws, block, workbook)
        return f"applied {len(op.blocks)} block(s) to sheet {op.sheet!r}"
    if op.kind == "add_sheet":
        if op.sheet.name in workbook.sheetnames:
            raise UnsupportedOperation(f"sheet already exists: {op.sheet.name!r}")
        ws = workbook.create_sheet(title=op.sheet.name)
        for block in op.sheet.blocks:
            _apply_block(ws, block, workbook)
        return f"added sheet {op.sheet.name!r}"
    raise UnsupportedOperation(f"unknown sheet op: {op.kind!r}")


def inspect_workbook(path: Path, *, max_rows_per_sheet: int = 40) -> dict[str, Any]:
    """Reopen a workbook: sheets, dimensions, formulas, charts, tables —
    and cached values for formula cells when the file carries them."""
    from openpyxl import load_workbook

    workbook = load_workbook(str(path))
    cached = load_workbook(str(path), data_only=True)
    sheets: list[dict[str, Any]] = []
    total_formulas = 0
    for ws in workbook.worksheets:
        formulas: list[str] = []
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    if len(formulas) < 20:
                        coordinate = cell.coordinate
                        formulas.append(coordinate)
                    total_formulas += 1
        cached_ws = cached[ws.title]
        computed: dict[str, Any] = {}
        for coordinate in formulas:
            value = cached_ws[coordinate].value
            if value is not None:
                computed[coordinate] = value
        sheets.append(
            {
                "name": ws.title,
                "dimensions": ws.calculate_dimension(),
                "charts": len(getattr(ws, "_charts", ())),
                "tables": sorted(ws.tables.keys()) if hasattr(ws, "tables") else [],
                "formula_cells": formulas,
                "formula_count": sum(
                    1
                    for row in ws.iter_rows()
                    for cell in row
                    if isinstance(cell.value, str) and cell.value.startswith("=")
                ),
                "cached_values": computed,
            }
        )
    return {
        "sheets": sheets,
        "sheet_names": workbook.sheetnames,
        "total_formulas": total_formulas,
    }


def recalculate(path: Path, out_path: Path, *, timeout_seconds: int, work_dir: Path) -> Path:
    """LibreOffice recalculation pass: convert xlsx→xlsx, embedding cached
    formula results while preserving the formulas. Returns the produced path."""
    import shutil

    from ..process import office_convert

    produced = office_convert(path, "xlsx", work_dir, timeout_seconds=timeout_seconds)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(str(produced), str(out_path))
    return out_path

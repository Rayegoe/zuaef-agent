"""Slides engine: declarative deck spec → themed .pptx via python-pptx,
bounded revisions by slide index, and structure inspection.

python-pptx is the single primary backend (spec pack 07/13): it creates AND
revises in one coherent engine, is pure Python (identical on x86_64 and
ARM64), and adds no Node runtime to any deployment. Layout defaults are
conservative and readable; the theme is ZUAEF-owned.
"""

from __future__ import annotations

from typing import Annotated, Any, Literal

from pydantic import BaseModel, Field, TypeAdapter, ValidationError

from .errors import UnsupportedOperation

MAX_SLIDES = 30
MAX_BULLETS = 8
MAX_BULLET_CHARS = 200
MAX_TITLE_CHARS = 120
MAX_TABLE_ROWS = 12
MAX_TABLE_COLS = 6

SLIDE_W_IN = 13.333  # 16:9
SLIDE_H_IN = 7.5

# ZUAEF-owned business themes (no copied template packs).
THEMES: dict[str, dict[str, str]] = {
    "business-light": {
        "background": "FFFFFF",
        "title": "1F3B63",
        "accent": "2E75B6",
        "text": "262626",
        "muted": "595959",
    },
    "business-dark": {
        "background": "1F2733",
        "title": "FFFFFF",
        "accent": "5B9BD5",
        "text": "E8EDF4",
        "muted": "AAB7C4",
    },
}


# ── deck spec ─────────────────────────────────────────────────────────────


class TitleSlide(BaseModel):
    kind: Literal["title"] = "title"


class TitleBodySlide(BaseModel):
    kind: Literal["title_body"] = "title_body"
    title: str
    bullets: list[str]

    def model_post_init(self, __context: Any, /) -> None:
        _validate_bullets(self.bullets)


class KeyMessageSlide(BaseModel):
    kind: Literal["key_message"] = "key_message"
    message: str

    def model_post_init(self, __context: Any, /) -> None:
        _bounded(self.message, 200, "message")


class TwoColumnSlide(BaseModel):
    kind: Literal["two_column"] = "two_column"
    title: str
    left_heading: str
    left_bullets: list[str]
    right_heading: str
    right_bullets: list[str]

    def model_post_init(self, __context: Any, /) -> None:
        _bounded(self.title, MAX_TITLE_CHARS, "title")
        _validate_bullets(self.left_bullets)
        _validate_bullets(self.right_bullets)


class ImageTextSlide(BaseModel):
    kind: Literal["image_text"] = "image_text"
    title: str
    image_path: str
    bullets: list[str]

    def model_post_init(self, __context: Any, /) -> None:
        _bounded(self.title, MAX_TITLE_CHARS, "title")
        _validate_bullets(self.bullets)


class TableSlide(BaseModel):
    kind: Literal["table"] = "table"
    title: str
    headers: list[str]
    rows: list[list[str]]

    def model_post_init(self, __context: Any, /) -> None:
        _bounded(self.title, MAX_TITLE_CHARS, "title")
        if not self.headers or len(self.headers) > MAX_TABLE_COLS:
            raise ValueError(f"table needs 1..{MAX_TABLE_COLS} columns")
        if len(self.rows) > MAX_TABLE_ROWS:
            raise ValueError(f"table exceeds {MAX_TABLE_ROWS} rows")
        for row in self.rows:
            if len(row) != len(self.headers):
                raise ValueError("every row must match the header width")
        for cell in [*self.headers, *[c for r in self.rows for c in r]]:
            # Cells may legitimately be empty; only bound length.
            if not isinstance(cell, str):
                raise TypeError("table cells must be strings")
            if len(cell) > 200:
                raise ValueError("table cell exceeds 200 characters")


class ChartSlide(BaseModel):
    kind: Literal["chart"] = "chart"
    title: str
    chart_type: Literal["bar", "line", "pie"] = "bar"
    categories: list[str]
    series_name: str
    values: list[float]

    def model_post_init(self, __context: Any, /) -> None:
        _bounded(self.title, MAX_TITLE_CHARS, "title")
        _bounded(self.series_name, 100, "series_name")
        if not self.categories or len(self.categories) > 12:
            raise ValueError("chart needs 1..12 categories")
        if len(self.values) != len(self.categories):
            raise ValueError("values must match categories")


class QuoteSlide(BaseModel):
    kind: Literal["quote"] = "quote"
    quote: str
    attribution: str | None = None

    def model_post_init(self, __context: Any, /) -> None:
        _bounded(self.quote, 400, "quote")


class ClosingSlide(BaseModel):
    kind: Literal["closing"] = "closing"
    title: str
    bullets: list[str] = Field(default_factory=list)

    def model_post_init(self, __context: Any, /) -> None:
        _bounded(self.title, MAX_TITLE_CHARS, "title")
        if self.bullets:
            _validate_bullets(self.bullets)


Slide = Annotated[
    TitleSlide | TitleBodySlide | KeyMessageSlide | TwoColumnSlide | ImageTextSlide | TableSlide | ChartSlide | QuoteSlide | ClosingSlide,
    Field(discriminator="kind"),
]


class DeckSpec(BaseModel):
    title: str
    subtitle: str | None = None
    theme: Literal["business-light", "business-dark"] = "business-light"
    slides: list[Slide]

    def model_post_init(self, __context: Any, /) -> None:
        _bounded(self.title, 150, "deck title")
        if not self.slides:
            raise ValueError("deck has no slides")
        if len(self.slides) > MAX_SLIDES:
            raise ValueError(f"deck exceeds {MAX_SLIDES} slides")


# ── revision ops ──────────────────────────────────────────────────────────


class ReplaceSlideTextOp(BaseModel):
    kind: Literal["replace_text"] = "replace_text"
    slide_index: int = Field(ge=1)
    find: str = Field(min_length=1, max_length=200)
    replacement: str


class SetSlideTitleOp(BaseModel):
    kind: Literal["set_title"] = "set_title"
    slide_index: int = Field(ge=1)
    title: str

    def model_post_init(self, __context: Any, /) -> None:
        _bounded(self.title, MAX_TITLE_CHARS, "title")


class SetSlideBulletsOp(BaseModel):
    kind: Literal["set_bullets"] = "set_bullets"
    slide_index: int = Field(ge=1)
    bullets: list[str]

    def model_post_init(self, __context: Any, /) -> None:
        _validate_bullets(self.bullets)


class UpdateSlideTableCellOp(BaseModel):
    kind: Literal["update_table_cell"] = "update_table_cell"
    slide_index: int = Field(ge=1)
    table_shape_index: int = Field(ge=0, default=0)
    row: int = Field(ge=0)
    col: int = Field(ge=0)
    text: str

    def model_post_init(self, __context: Any, /) -> None:
        _bounded(self.text, 200, "cell text")


class DeleteSlideOp(BaseModel):
    kind: Literal["delete_slide"] = "delete_slide"
    slide_index: int = Field(ge=1)


SlideOp = Annotated[
    ReplaceSlideTextOp | SetSlideTitleOp | SetSlideBulletsOp | UpdateSlideTableCellOp | DeleteSlideOp,
    Field(discriminator="kind"),
]


def _bounded(value: str, cap: int, what: str) -> None:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{what} must be a non-empty string")
    if len(value) > cap:
        raise ValueError(f"{what} exceeds {cap} characters")


def _validate_bullets(items: list[str]) -> None:
    if not items or len(items) > MAX_BULLETS:
        raise ValueError(f"bullets need 1..{MAX_BULLETS} items")
    for item in items:
        _bounded(item, MAX_BULLET_CHARS, "bullet")


def deck_from_json(data: dict) -> DeckSpec:
    try:
        return DeckSpec.model_validate(data)
    except ValidationError as exc:
        first = exc.errors()[0]
        location = ".".join(str(part) for part in first.get("loc", ()))
        raise ValueError(
            f"invalid deck spec at {location or '<root>'}: {first.get('msg')}"
        ) from exc


def slide_ops_from_json(data: list[dict]) -> list[SlideOp]:
    try:
        ops = TypeAdapter(list[SlideOp]).validate_python(data)
    except ValidationError as exc:
        first = exc.errors()[0]
        location = ".".join(str(part) for part in first.get("loc", ()))
        raise ValueError(
            f"invalid slide operation at {location or '<root>'}: {first.get('msg')}"
        ) from exc
    if not ops:
        raise ValueError("revision has no operations")
    if len(ops) > 30:
        raise ValueError("revision exceeds 30 operations")
    return ops


# ── build ─────────────────────────────────────────────────────────────────


def build_deck(spec: DeckSpec, workspace_root):
    """Render a DeckSpec into a python-pptx Presentation (unsaved)."""
    from pptx import Presentation
    from pptx.util import Inches

    theme = THEMES[spec.theme]
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)

    for index, slide in enumerate(spec.slides):
        if slide.kind == "title":
            new_slide = _blank(prs, theme)
            _add_title_slide(new_slide, spec.title, spec.subtitle, theme)
        else:
            new_slide = _blank(prs, theme)
            _add_content_slide(prs, new_slide, slide, theme, workspace_root)
    return prs


def _blank(prs, theme):
    """One blank slide with the theme background applied."""
    from pptx.dml.color import RGBColor

    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(theme["background"])
    return slide


def _add_title_slide(slide, title: str, subtitle: str | None, theme) -> None:
    from pptx.util import Inches

    box = slide.shapes.add_textbox(Inches(0.9), Inches(2.6), Inches(11.5), Inches(1.4))
    _style_text(box.text_frame.paragraphs[0], title, size=40, bold=True, color=theme["title"])
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.9), Inches(4.1), Inches(11.5), Inches(0.8))
        _style_text(sub.text_frame.paragraphs[0], subtitle, size=18, color=theme["muted"])


def _add_content_slide(prs, slide, slide_spec, theme, workspace_root) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches

    kind = slide_spec.kind

    title_text = getattr(slide_spec, "title", None)
    if kind == "key_message":
        title_text = None

    if title_text:
        box = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(0.9))
        _style_text(
            box.text_frame.paragraphs[0], title_text, size=28, bold=True,
            color=theme["title"],
        )

    if kind == "key_message":
        box = slide.shapes.add_textbox(Inches(1.2), Inches(2.8), Inches(10.9), Inches(2))
        paragraph = box.text_frame.paragraphs[0]
        _style_text(paragraph, slide_spec.message, size=32, bold=True, color=theme["accent"])
        paragraph.alignment = PP_ALIGN.CENTER
        return

    if kind == "title_body":
        _add_bullets(slide, Inches(0.9), Inches(1.7), Inches(11.5), Inches(5), slide_spec.bullets, theme)
        return

    if kind == "two_column":
        left = slide.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(5.9), Inches(5.2))
        _style_text(left.text_frame.paragraphs[0], slide_spec.left_heading, size=18, bold=True, color=theme["accent"])
        _add_bullets(slide, Inches(0.9), Inches(2.3), Inches(5.6), Inches(4.4), slide_spec.left_bullets, theme)
        right = slide.shapes.add_textbox(Inches(6.9), Inches(1.6), Inches(5.9), Inches(5.2))
        _style_text(right.text_frame.paragraphs[0], slide_spec.right_heading, size=18, bold=True, color=theme["accent"])
        _add_bullets(slide, Inches(7.1), Inches(2.3), Inches(5.6), Inches(4.4), slide_spec.right_bullets, theme)
        return

    if kind == "image_text":
        from ..paths import resolve_input_path

        source = resolve_input_path(workspace_root, slide_spec.image_path, max_bytes=25_000_000)
        slide.shapes.add_picture(str(source), Inches(0.7), Inches(1.9), width=Inches(6.2))
        _add_bullets(slide, Inches(7.3), Inches(1.9), Inches(5.4), Inches(4.8), slide_spec.bullets, theme)
        return

    if kind == "table":
        rows, cols = 1 + len(slide_spec.rows), len(slide_spec.headers)
        shape = slide.shapes.add_table(rows, cols, Inches(0.7), Inches(1.8), Inches(11.9), Inches(0.5 * rows))
        table = shape.table
        for col, header in enumerate(slide_spec.headers):
            table.cell(0, col).text = str(header)
        for row_i, row in enumerate(slide_spec.rows, start=1):
            for col, value in enumerate(row):
                table.cell(row_i, col).text = str(value)
        return

    if kind == "chart":
        data = CategoryChartData()
        data.categories = [str(c) for c in slide_spec.categories]
        data.add_series(slide_spec.series_name, tuple(float(v) for v in slide_spec.values))
        chart_type = {
            "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "line": XL_CHART_TYPE.LINE_MARKERS,
            "pie": XL_CHART_TYPE.PIE,
        }[slide_spec.chart_type]
        frame = slide.shapes.add_chart(
            chart_type, Inches(1.2), Inches(1.7), Inches(10.9), Inches(5.2), data
        )
        frame.chart.has_legend = slide_spec.chart_type == "pie"
        return

    if kind == "quote":
        box = slide.shapes.add_textbox(Inches(1.4), Inches(2.4), Inches(10.5), Inches(1.8))
        paragraph = box.text_frame.paragraphs[0]
        _style_text(paragraph, f"“{slide_spec.quote}”", size=26, italic=True, color=theme["text"])
        paragraph.alignment = PP_ALIGN.CENTER
        if slide_spec.attribution:
            attr = slide.shapes.add_textbox(Inches(1.4), Inches(4.5), Inches(10.5), Inches(0.6))
            attr_paragraph = attr.text_frame.paragraphs[0]
            _style_text(attr_paragraph, f"— {slide_spec.attribution}", size=14, color=theme["muted"])
            attr_paragraph.alignment = PP_ALIGN.CENTER
        return

    if kind == "closing":
        box = slide.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.9), Inches(1.2))
        paragraph = box.text_frame.paragraphs[0]
        _style_text(paragraph, slide_spec.title, size=34, bold=True, color=theme["title"])
        paragraph.alignment = PP_ALIGN.CENTER
        if slide_spec.bullets:
            _add_bullets(slide, Inches(2.4), Inches(3.8), Inches(8.5), Inches(2.8), slide_spec.bullets, theme, centered=True)
        return

    raise UnsupportedOperation(f"unknown slide kind: {kind!r}")


def _add_bullets(slide, left, top, width, height, items, theme, *, centered: bool = False) -> None:
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt

    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    for i, item in enumerate(items):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        _style_text(paragraph, f"• {item}", size=16, color=theme["text"])
        paragraph.space_after = Pt(8)
        if centered:
            paragraph.alignment = PP_ALIGN.CENTER


def _style_text(paragraph, text: str, *, size: int, color: str, bold: bool = False, italic: bool = False):
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    paragraph.text = text
    for run in paragraph.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = RGBColor.from_string(color)


# ── bounded revisions ─────────────────────────────────────────────────────


def _slide_at(prs, index_1based: int):
    if index_1based > len(prs.slides):
        raise UnsupportedOperation(
            f"slide {index_1based} out of range ({len(prs.slides)} slides)"
        )
    return prs.slides[index_1based - 1]


def _iter_text_shapes(slide):
    for shape in slide.shapes:
        if shape.has_text_frame:
            yield shape


def _rewrite_paragraph(paragraph, new_text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = new_text


def apply_slide_op(prs, op: Any, workspace_root) -> str:
    kind = op.kind

    if kind == "replace_text":
        slide = _slide_at(prs, op.slide_index)
        for shape in _iter_text_shapes(slide):
            for paragraph in shape.text_frame.paragraphs:
                if op.find in paragraph.text:
                    new_text = paragraph.text.replace(op.find, op.replacement, 1)
                    _rewrite_paragraph(paragraph, new_text)
                    return f"replaced text on slide {op.slide_index}"
        raise UnsupportedOperation(f"find text not found on slide {op.slide_index}: {op.find!r}")

    if kind == "set_title":
        slide = _slide_at(prs, op.slide_index)
        shapes = list(_iter_text_shapes(slide))
        if not shapes:
            raise UnsupportedOperation(f"slide {op.slide_index} has no text shape to hold a title")
        _rewrite_paragraph(shapes[0].text_frame.paragraphs[0], op.title)
        return f"set title of slide {op.slide_index}"

    if kind == "set_bullets":
        from pptx.util import Pt

        slide = _slide_at(prs, op.slide_index)
        shapes = [
            shape
            for shape in _iter_text_shapes(slide)
            if len(shape.text_frame.paragraphs) > 1
            or any(p.text.startswith("•") for p in shape.text_frame.paragraphs)
        ]
        if not shapes:
            raise UnsupportedOperation(
                f"slide {op.slide_index} has no bullet body to replace"
            )
        body = max(shapes, key=lambda s: s.width * s.height)
        frame = body.text_frame
        frame.clear()
        for i, item in enumerate(op.bullets):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            _rewrite_paragraph(paragraph, f"• {item}")
            for run in paragraph.runs:
                run.font.size = Pt(16)
        return f"set {len(op.bullets)} bullet(s) on slide {op.slide_index}"

    if kind == "update_table_cell":
        slide = _slide_at(prs, op.slide_index)
        tables = [shape for shape in slide.shapes if getattr(shape, "has_table", False)]
        if op.table_shape_index >= len(tables):
            raise UnsupportedOperation(
                f"slide {op.slide_index} has no table shape {op.table_shape_index}"
            )
        table = tables[op.table_shape_index].table
        try:
            table.cell(op.row, op.col).text = op.text
        except IndexError as exc:
            raise UnsupportedOperation(
                f"cell ({op.row}, {op.col}) out of range on slide {op.slide_index}"
            ) from exc
        return f"updated slide {op.slide_index} table cell ({op.row}, {op.col})"

    if kind == "delete_slide":
        _slide_at(prs, op.slide_index)
        xml_slides = prs.slides._sldIdLst
        slide_ids = list(xml_slides)
        xml_slides.remove(slide_ids[op.slide_index - 1])
        return f"deleted slide {op.slide_index}"

    raise UnsupportedOperation(f"unknown slide op: {kind!r}")


# ── inspection ────────────────────────────────────────────────────────────


def inspect_deck(path) -> dict[str, Any]:
    from pptx import Presentation

    prs = Presentation(str(path))
    slides: list[dict[str, Any]] = []
    for index, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        tables = 0
        charts = 0
        pictures = 0
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                tables += 1
                continue
            if getattr(shape, "has_chart", False):
                charts += 1
                continue
            if shape.shape_type is not None and "PICTURE" in str(shape.shape_type):
                pictures += 1
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip()[:160])
        slides.append(
            {
                "slide": index,
                "texts": texts[:6],
                "tables": tables,
                "charts": charts,
                "pictures": pictures,
            }
        )
    return {"slides": slides, "count": len(slides)}

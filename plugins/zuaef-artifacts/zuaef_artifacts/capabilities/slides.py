"""Slides artifact capability: create / revise / inspect presentation decks
under ``workspace/artifacts/slides/``."""

from __future__ import annotations

import time
from pathlib import Path
from typing import Any

from pydantic_ai import FunctionToolset
from pydantic_ai.capabilities import Capability

from zuaef_agent.models import CoreDeps

from ..contracts import ArtifactBounds, ArtifactResult, failed_result
from ..engines import slides_engine
from ..engines.errors import UnsupportedOperation
from ..guidance import load_guidance
from ..paths import (
    PathRejected,
    allocate_output_path,
    allocate_render_dir,
    allocate_revision_path,
    allocate_work_dir,
    artifact_display_path,
    cleanup_dir,
    resolve_existing_artifact,
)
from ..process import MissingDependency, ProcessFailed
from ..qa import qa_slides

DESCRIPTION = (
    "Presentation-ready slide decks (.pptx) 给老板汇报/周会材料/客户pitch/"
    "项目复盘/评审汇报/strategy briefing: management reviews, client pitches, "
    "project updates, weekly/monthly reviews. Use when the outcome is a "
    "visual briefing someone will present or project, or when the user asks "
    "for PPT/slides."
)

_ERROR_CODES: tuple[tuple[type[Exception], str], ...] = (
    (PathRejected, "path_rejected"),
    (UnsupportedOperation, "unsupported_operation"),
    (MissingDependency, "missing_dependency"),
    (ProcessFailed, "conversion_failed"),
)


def _error_code_for(exc: Exception) -> str:
    for exc_type, code in _ERROR_CODES:
        if isinstance(exc, exc_type):
            return code
    return "invalid_request" if isinstance(exc, ValueError) else "internal_error"


def _default_stem() -> str:
    return f"deck-{time.strftime('%Y%m%d-%H%M%S')}"


def build_capability(
    *,
    bounds: ArtifactBounds,
    workspace_root: Path,
    work_root: Path,
) -> Capability[CoreDeps]:
    toolset: FunctionToolset[CoreDeps] = FunctionToolset[CoreDeps]()

    @toolset.tool_plain
    def create_slides(
        spec: slides_engine.DeckSpec, output_name: str | None = None
    ) -> ArtifactResult:
        """Create a presentation deck (.pptx) in one call. spec: deck title,
        subtitle, theme (business-light | business-dark), and slides[] —
        kinds: title; title_body(title, bullets); key_message(message);
        two_column(title, left_heading, left_bullets, right_heading,
        right_bullets); image_text(title, image_path, bullets); table(title,
        headers, rows); chart(title, chart_type bar|line|pie, categories,
        series_name, values); quote(quote, attribution?); closing(title,
        bullets?). Max 30 slides; one idea per slide. Runs reopen/convert/
        render QA inside this call."""
        out_path: Path | None = None
        work: Path | None = None
        render: Path | None = None
        try:
            out_path = allocate_output_path(
                workspace_root, "slides", output_name, default_stem=_default_stem()
            )
            prs = slides_engine.build_deck(spec, workspace_root)
            prs.save(str(out_path))

            work = allocate_work_dir(work_root, "slides-create")
            render = allocate_render_dir(work_root, "slides-create")
            outcome = qa_slides(
                out_path,
                expected_slides=len(spec.slides),
                work_dir=work,
                render_dir=render,
                bounds=bounds,
            )
            if outcome.qa_state == "failed":
                out_path.unlink(missing_ok=True)
                return ArtifactResult(
                    ok=False,
                    artifact_type="slides",
                    qa_state="failed",
                    error_code="qa_failed",
                    warnings=outcome.warnings,
                    summary="deck was created but failed QA; it was removed",
                )
            return ArtifactResult(
                ok=True,
                artifact_path=artifact_display_path(workspace_root, out_path),
                artifact_type="slides",
                qa_state=outcome.qa_state,  # type: ignore[arg-type]
                warnings=outcome.warnings,
                summary=(
                    f"created deck '{spec.title}' with {len(spec.slides)} slide(s) "
                    f"({spec.theme} theme)"
                ),
            )
        except Exception as exc:  # noqa: BLE001 - recoverable, format-local
            if out_path is not None:
                out_path.unlink(missing_ok=True)
            return failed_result("slides", _error_code_for(exc), exc)  # type: ignore[arg-type]
        finally:
            if work is not None:
                cleanup_dir(work)
            if render is not None:
                cleanup_dir(render)

    @toolset.tool_plain
    def revise_slides(
        path: str,
        operations: list[slides_engine.SlideOp],
        output_name: str | None = None,
    ) -> ArtifactResult:
        """Apply bounded edits to an existing deck by slide index (1-based).
        Operations: replace_text(slide_index, find, replacement),
        set_title(slide_index, title), set_bullets(slide_index, bullets),
        update_table_cell(slide_index, row, col, text), delete_slide(
        slide_index). Writes a versioned copy (…-rev2.pptx); the source file
        is never modified in place. QA re-runs inside this call."""
        out_path: Path | None = None
        work: Path | None = None
        render: Path | None = None
        try:
            source = resolve_existing_artifact(
                workspace_root, path, "slides", max_bytes=bounds.max_input_bytes
            )
            out_path = allocate_revision_path(
                workspace_root, "slides", source, output_name
            )
            from pptx import Presentation

            prs = Presentation(str(source))
            changes = [slides_engine.apply_slide_op(prs, op, workspace_root) for op in operations]
            prs.save(str(out_path))

            work = allocate_work_dir(work_root, "slides-revise")
            render = allocate_render_dir(work_root, "slides-revise")
            outcome = qa_slides(
                out_path,
                expected_slides=len(prs.slides),
                work_dir=work,
                render_dir=render,
                bounds=bounds,
            )
            if outcome.qa_state == "failed":
                out_path.unlink(missing_ok=True)
                return ArtifactResult(
                    ok=False,
                    artifact_type="slides",
                    qa_state="failed",
                    error_code="qa_failed",
                    warnings=outcome.warnings,
                    summary="revision was applied but QA failed; it was removed",
                )
            return ArtifactResult(
                ok=True,
                artifact_path=artifact_display_path(workspace_root, out_path),
                artifact_type="slides",
                qa_state=outcome.qa_state,  # type: ignore[arg-type]
                warnings=outcome.warnings,
                summary=f"revised '{source.name}': {'; '.join(changes)}",
            )
        except Exception as exc:  # noqa: BLE001 - recoverable, format-local
            if out_path is not None:
                out_path.unlink(missing_ok=True)
            return failed_result("slides", _error_code_for(exc), exc)  # type: ignore[arg-type]
        finally:
            if work is not None:
                cleanup_dir(work)
            if render is not None:
                cleanup_dir(render)

    @toolset.tool_plain
    def inspect_slides(path: str) -> dict[str, Any]:
        """Inspect an existing deck: per-slide texts and table/chart/picture
        counts, plus total slide count."""
        try:
            source = resolve_existing_artifact(
                workspace_root, path, "slides", max_bytes=bounds.max_input_bytes
            )
            facts = slides_engine.inspect_deck(source)
            facts.update(
                {
                    "path": artifact_display_path(workspace_root, source),
                    "bytes": source.stat().st_size,
                }
            )
            return facts
        except Exception as exc:  # noqa: BLE001 - recoverable, format-local
            return {"error": _error_code_for(exc), "message": str(exc)[:500], "path": path}

    return Capability[CoreDeps](
        id="slides-artifacts",
        description=DESCRIPTION,
        instructions=load_guidance("slides.md"),
        toolsets=[toolset],
        defer_loading=True,
    )

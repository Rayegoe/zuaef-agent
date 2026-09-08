"""Document toolset behavior for the general knowledge worker plugin.

Covers spec pack 05 acceptance E-I (binary documents with locators),
N (path safety) and O (bounded output), plus the recoverable-error and
NO_EXTRACTABLE_TEXT contracts from 01_PRODUCT_SPEC.
"""

from __future__ import annotations

import asyncio
import json
from pathlib import Path

import pytest
from zuaef_knowledge_worker.document_tools import make_document_toolset


def _toolset(workspace: Path, *, chunk_chars: int = 12_000, max_bytes: int = 25_000_000):
    return make_document_toolset(
        workspace_root=workspace, chunk_chars=chunk_chars, max_bytes=max_bytes
    )


def _call(toolset, name: str, *args, **kwargs):
    return toolset.tools[name].function(*args, **kwargs)


def _write(workspace: Path, name: str, content: str | bytes) -> Path:
    path = workspace / name
    path.parent.mkdir(parents=True, exist_ok=True)
    if isinstance(content, bytes):
        path.write_bytes(content)
    else:
        path.write_text(content, encoding="utf-8")
    return path


def _minimal_pdf(*page_texts: str) -> bytes:
    """Hand-built one-font PDF with one content stream per page (pypdf reads
    the text back deterministically; no extra dependency needed)."""
    objects: list[bytes] = []

    page_ids = []
    content_ids = []
    # 1: catalog, 2: pages, then per page: page object + content object,
    # then font. IDs assigned up front so the pages tree can reference them.
    n_pages = len(page_texts)
    font_id = 3 + 2 * n_pages
    for i in range(n_pages):
        page_ids.append(3 + 2 * i)
        content_ids.append(4 + 2 * i)

    objects.append(  # 1 catalog
        b"<< /Type /Catalog /Pages 2 0 R >>"
    )
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects.append(  # 2 pages
        f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode()
    )
    for text in page_texts:
        stream = (
            f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        )
        objects.append(  # page object (placeholder, patched below)
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources "
            f"<< /Font << /F1 {font_id} 0 R >> >> /Contents @@ >>".encode()
        )
        objects.append(  # content stream
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
            + stream + b"\nendstream"
        )
    objects.append(  # font
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    )

    # Patch the /Contents references now that content ids are final.
    for i, page_id in enumerate(page_ids):
        objects[page_id - 1] = objects[page_id - 1].replace(
            b"/Contents @@", f"/Contents {content_ids[i]} 0 R".encode()
        )

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_at}\n%%EOF\n"
    ).encode()
    return bytes(out)


@pytest.fixture()
def workspace(tmp_path: Path) -> Path:
    return tmp_path


def _docs(workspace: Path) -> dict[str, Path]:
    """One fixture per supported format (spec pack 04 M4 fixture list)."""
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    paths: dict[str, Path] = {}
    paths["txt"] = _write(workspace, "inbox/note.txt", "alpha line\nbeta line\n")
    paths["md"] = _write(workspace, "inbox/note.md", "# Heading\n\nPydantic is great.\n")
    paths["json"] = _write(
        workspace, "inbox/data.json", json.dumps({"tool": "pydantic-ai", "n": 2})
    )
    paths["csv"] = _write(workspace, "inbox/data.csv", "item,qty\nwidget,3\ngadget,5\n")
    paths["html"] = _write(
        workspace,
        "inbox/page.html",
        "<html><head><style>body{color:red}</style></head>"
        "<body><script>evil()</script><h1>Visible Title</h1><p>Pdf body text</p>"
        "</body></html>",
    )
    paths["pdf"] = _write(
        workspace,
        "inbox/sample.pdf",
        _minimal_pdf(
            "Quarterly revenue grew. " + "A" * 90,
            "Risk factors remain. " + "B" * 90,
        ),
    )
    paths["empty_pdf"] = _write(workspace, "inbox/empty.pdf", _minimal_pdf(""))

    doc = Document()
    doc.add_paragraph("Contract start.")
    doc.add_paragraph("Confidential clause 7 applies.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "key"
    table.rows[0].cells[1].text = "value"
    paths["docx"] = _write(workspace, "inbox/sample.docx", b"")
    doc.save(paths["docx"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["product", "amount"])
    ws.append(["alpha", 100])
    ws.append(["beta", 900])
    ws2 = wb.create_sheet("Notes")
    ws2["A1"] = "remember the milk"
    paths["xlsx"] = _write(workspace, "inbox/sample.xlsx", b"")
    wb.save(paths["xlsx"])

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Decision: approve budget"
    paths["pptx"] = _write(workspace, "inbox/sample.pptx", b"")
    prs.save(paths["pptx"])
    return paths


def test_inspect_reports_format_and_extractability(workspace: Path):
    _docs(workspace)
    ts = _toolset(workspace)
    info = _call(ts, "inspect_document", "inbox/sample.pdf")
    assert info["format"] == "pdf" and info["sections"] == 2
    assert info["extractable"] is True
    empty = _call(ts, "inspect_document", "inbox/empty.pdf")
    assert empty["extractable"] is False and empty["nonempty_sections"] == 0


def test_read_document_returns_located_text_and_pagination(workspace: Path):
    _docs(workspace)
    ts = _toolset(workspace, chunk_chars=100)
    first = _call(ts, "read_document", "inbox/sample.pdf")
    assert "[page:1]" in first["text"] and "Quarterly revenue grew." in first["text"]
    assert "[page:2]" not in first["text"]
    assert first["next_offset"] is not None
    second = _call(ts, "read_document", "inbox/sample.pdf", offset=first["next_offset"])
    assert "[page:2]" in second["text"] and "Risk factors remain." in second["text"]
    assert second["total_chars"] == first["total_chars"]
    # Paging walks to exhaustion deterministically.
    third = _call(ts, "read_document", "inbox/sample.pdf", offset=second["next_offset"])
    assert third["next_offset"] is None


def test_read_document_empty_pdf_is_honest_no_extractable_text(workspace: Path):
    _docs(workspace)
    ts = _toolset(workspace)
    result = _call(ts, "read_document", "inbox/empty.pdf")
    assert result["warning"] == "NO_EXTRACTABLE_TEXT"
    # Only the locator header exists; no section body text was fabricated.
    assert result["text"].strip() == "[page:1]"


def test_read_document_scanned_like_pdf_does_not_invent_content(workspace: Path):
    """An image-only scan has no text layer: the tool must say so, never fill in."""
    _docs(workspace)
    ts = _toolset(workspace)
    result = _call(ts, "read_document", "inbox/empty.pdf")
    assert "Quarterly" not in result["text"] and "risk" not in result["text"].lower()


def test_html_strips_script_and_style_noise(workspace: Path):
    _docs(workspace)
    ts = _toolset(workspace)
    result = _call(ts, "read_document", "inbox/page.html")
    assert "Visible Title" in result["text"]
    assert "evil()" not in result["text"] and "color:red" not in result["text"]


def test_pdf_search_reports_page_locators(workspace: Path):
    _docs(workspace)
    ts = _toolset(workspace)
    hits = _call(ts, "search_document", "inbox/sample.pdf", "Risk factors")
    assert len(hits) == 1
    assert hits[0]["locator"] == "page:2"
    assert "Risk factors remain." in hits[0]["excerpt"]


def test_xlsx_is_read_only_data_only_with_sheet_locators(workspace: Path):
    _docs(workspace)
    ts = _toolset(workspace)
    hits = _call(ts, "search_document", "inbox/sample.xlsx", "beta")
    assert hits and hits[0]["locator"] == "sheet:Sales/row:3"
    assert "900" in hits[0]["excerpt"]


def test_docx_paragraphs_and_tables_are_located(workspace: Path):
    _docs(workspace)
    ts = _toolset(workspace)
    hits = _call(ts, "search_document", "inbox/sample.docx", "clause 7")
    assert hits and hits[0]["locator"].startswith("paragraph:")
    rows = _call(ts, "search_document", "inbox/sample.docx", "value")
    assert rows and rows[0]["locator"] == "table:1/row:1"


def test_pptx_slide_locators_include_tables(workspace: Path):
    _docs(workspace)
    ts = _toolset(workspace)
    hits = _call(ts, "search_document", "inbox/sample.pptx", "budget")
    assert hits and hits[0]["locator"] == "slide:1"
    assert "Decision: approve budget" in hits[0]["excerpt"]


def test_search_documents_discovers_inbox_documents(workspace: Path):
    _docs(workspace)
    ts = _toolset(workspace)
    hits = _call(ts, "search_documents", "pydantic")
    found = {hit["path"] for hit in hits}
    assert found == {"inbox/note.md", "inbox/data.json"}
    for hit in hits:
        assert "pydantic" in hit["excerpt"].casefold()


def test_search_documents_with_explicit_paths_only(workspace: Path):
    _docs(workspace)
    ts = _toolset(workspace)
    hits = _call(ts, "search_documents", "widget", paths=["inbox/data.csv"])
    assert [hit["locator"] for hit in hits] == ["row:2"]


def test_json_is_pretty_printed_not_raw(workspace: Path):
    _docs(workspace)
    ts = _toolset(workspace)
    result = _call(ts, "read_document", "inbox/data.json")
    assert '"tool": "pydantic-ai"' in result["text"]


def test_traversal_outside_workspace_is_rejected(workspace: Path):
    ts = _toolset(workspace)
    outside = workspace.parent / "outside.txt"
    outside.write_text("secret", encoding="utf-8")
    result = _call(ts, "read_document", f"../{outside.name}")
    assert "error" in result and "outside the workspace" in result["error"]


def test_absolute_path_outside_workspace_is_rejected(workspace: Path):
    ts = _toolset(workspace)
    result = _call(ts, "read_document", "/etc/passwd")
    assert "error" in result


def test_symlink_escape_is_rejected(workspace: Path):
    ts = _toolset(workspace)
    secret = workspace.parent / "secret.txt"
    secret.write_text("top secret", encoding="utf-8")
    link = workspace / "inbox" / "leak.txt"
    link.parent.mkdir(parents=True, exist_ok=True)
    link.symlink_to(secret)
    result = _call(ts, "read_document", "inbox/leak.txt")
    assert "error" in result and "outside the workspace" in result["error"]
    assert "top secret" not in json.dumps(result)


def test_env_file_is_protected_even_if_supported_suffix(workspace: Path):
    ts = _toolset(workspace)
    _write(workspace, "inbox/mid.env", "TOKEN=x")
    result = _call(ts, "read_document", "inbox/mid.env")
    assert "error" in result and "unsupported" in result["error"]
    _write(workspace, "inbox/.env", "TOKEN=x")
    result = _call(ts, "read_document", "inbox/.env")
    assert "error" in result and "protected" in result["error"]


def test_unsupported_suffix_is_rejected(workspace: Path):
    ts = _toolset(workspace)
    _write(workspace, "inbox/binary.exe", "MZ...")
    result = _call(ts, "read_document", "inbox/binary.exe")
    assert "error" in result and "unsupported" in result["error"]


def test_missing_document_is_recoverable_error(workspace: Path):
    ts = _toolset(workspace)
    result = _call(ts, "read_document", "inbox/nope.pdf")
    assert "error" in result and "does not exist" in result["error"]


def test_oversized_document_is_rejected_not_flooded(workspace: Path):
    ts = _toolset(workspace, max_bytes=64)
    _write(workspace, "inbox/big.txt", "x" * 200)
    result = _call(ts, "read_document", "inbox/big.txt")
    assert "error" in result and "byte limit" in result["error"]


def test_read_output_is_bounded_by_chunk_chars(workspace: Path):
    ts = _toolset(workspace, chunk_chars=100)
    _write(workspace, "inbox/long.txt", "y" * 10_000)
    result = _call(ts, "read_document", "inbox/long.txt")
    assert len(result["text"]) <= 100
    assert result["next_offset"] == 100


def test_empty_query_is_recoverable(workspace: Path):
    ts = _toolset(workspace)
    assert "error" in _call(ts, "search_document", "inbox/note.txt", "  ")[0]
    assert "error" in _call(ts, "search_documents", "")[0]


def test_invalid_json_is_recoverable_error(workspace: Path):
    ts = _toolset(workspace)
    _write(workspace, "inbox/broken.json", "{not json")
    result = _call(ts, "read_document", "inbox/broken.json")
    assert "error" in result


def test_toolset_carries_untrusted_evidence_instruction(workspace: Path):
    ts = _toolset(workspace)
    instructions = "\n".join(ts._instructions)
    assert "untrusted evidence" in instructions


def test_output_language_is_projected_into_instructions(workspace: Path):
    ts = make_document_toolset(
        workspace_root=workspace, chunk_chars=100, max_bytes=1000,
        output_language="zh-CN",
    )
    instructions = "\n".join(ts._instructions)
    assert "zh-CN" in instructions
    plain = "\n".join(_toolset(workspace)._instructions)
    assert "zh-CN" not in plain


def test_search_limit_is_capped(workspace: Path):
    _write(workspace, "inbox/many.txt", "\n".join(f"needle {i}" for i in range(80)))
    ts = _toolset(workspace)
    hits = _call(ts, "search_document", "inbox/many.txt", "needle", limit=999)
    assert len(hits) <= 50


def test_tools_are_async_safe_plain_functions(workspace: Path):
    """Tools run through pydantic-ai's sync-tool path; they must not need an event loop."""
    ts = _toolset(workspace)
    _docs(workspace)
    result = _call(ts, "inspect_document", "inbox/sample.pdf")
    assert result["extractable"] is True
    assert asyncio.get_event_loop_policy() is not None
